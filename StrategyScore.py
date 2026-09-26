import openpyxl
import os
import shutil

#################################### CHANGE ME #####################################

#SET DIRECTORY NAMES
#must be subdirectories of the main directory the program resides in

#INPUT DIRECTORY: for the directory that holds the formatted excel sheets to code:
input_directory = 'Input Folder'

#OUTPUT DIRECTORY: for the directory to place the newly-coded excel sheets:
output_directory = 'Output Folder'

#SET DICTIONARY NAME: for the excel workbook containing the dictionary of words for each category:
dictionary_name = 'Dictionary.xlsx'

######################################################################################

#global variables
deter_list = []
norm_list = []
pe_list = []

#FUNCTION DEFINITIONS
#checks a given word against a dictionary list and codes that row, returning frequency if a match is found
def code_cell(row, col, number, color):
    #set flag for detecting a partial match to false
    partial_word_flag = False
    #identify cells
    word_cell = ws.cell(row=row, column=col)
    frequency_cell = ws.cell(row=row, column=col+1)
    code_cell = ws.cell(row=row, column=col+2)
    comment_cell = ws.cell(row=row, column=col+3)

    #determine which category list to use by input number
    if number == 1:
        list = deter_list
    elif number == 2:
        list = norm_list
    elif number == 3:
        list = pe_list
    else:
        return 0
    
    #for each entry in the dictionary, check it against the given word
    for item in list:
            if (item == word_cell.value):
                #for an exact match, code number & color parameter and return frequency right away
                code_cell.value = number
                code_cell.fill = openpyxl.styles.fills.PatternFill(start_color=color, end_color=color, fill_type="solid")
                print('found')
                return frequency_cell.value
            elif (item in word_cell.value):
                #checks if the dictionary word/phrase is in the cell
                #and if it's not the first character,
                #that the previous character is not a letter (word inside a word)
                if (word_cell.value.index(item) == 0) or (not(word_cell.value[(word_cell.value.index(item))-1].isalpha())):
                    #for a partial match, code number & color and set partial match flag to true
                    code_cell.value = number
                    code_cell.fill = openpyxl.styles.fills.PatternFill(start_color=color, end_color=color, fill_type="solid")
                    partial_word_flag = True
                    partial_word_index = list.index(item)
                    print('partial found')
    
    #for a partial match flag, mark the comment column in orange and return the frequency
    if partial_word_flag:
        comment_cell.fill = openpyxl.styles.fills.PatternFill(start_color="FFE6B3", end_color="FFE6B3", fill_type="solid")
        comment_cell.value = "from '" + list[partial_word_index] + "'"

        return frequency_cell.value
    else:
        #no match, return the frequency as 0
        return 0

#code and calculate for a given column of words/phrases based on format
def code_column(col):
    #set variables
    curr_row = 2
    total_deter = 0
    total_norm = 0
    total_pe = 0
    total_count = 0


    #iterate through phrases - each for loop checks if the phrase contains a key word from the list
    #if it does, it adds the frequency to the frequency total for its category, and codes the appropriate column with a number and colour
    while (ws.cell(row=curr_row, column=col).value != None):
        #fill with default 0
        ws.cell(row=curr_row, column=col+2).value = 0

        #set count of repeat assigned categories to 0, then run through the three categories
        repeat_categories = 0

        for i in range(0,3):
            if code_cell(curr_row, col, i+1, "DDFF99") > 0:
                repeat_categories = repeat_categories + 1

        #code appropriate number & colors, adding the frquency number to appropriate total
        total_deter = total_deter + code_cell(curr_row, col, 1, "E6FFCC")

        total_norm = total_norm + code_cell(curr_row, col, 2, "99CCFF")

        total_pe = total_pe + code_cell(curr_row, col, 3, "B3B3FF")

        #if the cell contains words from multiple categories (i.e. repeat_categories > 1), give warning
        if repeat_categories > 1: 
            ws.cell(row=curr_row, column = col+3).fill = openpyxl.styles.fills.PatternFill(start_color="FF8080", end_color="FF8080", fill_type="solid")
            ws.cell(row=curr_row, column = col+3).value = "contains CONFLICTING dictionary categories"

        #add frequency to general total
        total_count = total_count + ws.cell(row=curr_row, column=col+1).value

        #increment the row
        curr_row = curr_row + 1

    #calculate block
    ws.cell(row=2, column=col+6).value = total_deter / total_count
    ws.cell(row=3, column=col+6).value = total_norm / total_count
    ws.cell(row=4, column=col+6).value = total_pe / total_count
    ws.cell(row=6, column=col+6).value = ws.cell(row=2, column=col+6).value + ws.cell(row=3, column=col+6).value + ws.cell(row=4, column=col+6).value
    
    ws.cell(row=2, column=col+7).value = total_deter / (total_deter + total_norm + total_pe)
    ws.cell(row=3, column=col+7).value = total_norm / (total_deter + total_norm + total_pe)
    ws.cell(row=4, column=col+7).value = total_pe / (total_deter + total_norm + total_pe)
    ws.cell(row=6, column=col+7).value = ws.cell(row=2, column=col+7).value + ws.cell(row=3, column=col+7).value + ws.cell(row=4, column=col+7).value

    

#gathers dictionary reference lists from dictionary excel sheet
def gather_dictionary_lists():
    #open dictionary workbook & sheet
    dictionary = openpyxl.load_workbook(filename = dictionary_name, read_only = True)
    dic_ws = dictionary.active

    #create dictionary lists (deterrence, norms, PE)

    #deterrence
    dic_row = 2
    while (dic_ws.cell(row= dic_row, column=1).value != None):
        deter_list.append(dic_ws.cell(dic_row, column=1).value.upper())
        dic_row = dic_row + 1

    #norms
    dic_row = 2
    while (dic_ws.cell(row= dic_row, column=3).value != None):
        norm_list.append(dic_ws.cell(dic_row, column=3).value.upper())
        dic_row = dic_row + 1

    #PE
    dic_row = 2
    while (dic_ws.cell(row= dic_row, column=5).value != None):
        pe_list.append(dic_ws.cell(dic_row, column=5).value.upper())
        dic_row = dic_row + 1


#MAIN
#create dictionary lists for reference
gather_dictionary_lists()

#iterate through input directory and code each workbook, then save it
for entry in os.scandir(input_directory):
    if entry.name.endswith(".xlsx"):
        wb = openpyxl.load_workbook(filename = entry)
        ws = wb.active

        code_column(1)
        code_column(10)
        code_column(19)

        wb.save('new_' + entry.name)

#move the coded workbooks into the output directory
for entry in os.scandir('.'):
    if entry.name.endswith('.xlsx') and entry.name.startswith('new_'):
        shutil.move(entry, output_directory)
