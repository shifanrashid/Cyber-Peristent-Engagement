import openpyxl
import os
import xlwings as xw


############### CHANGE ME #####################

#Name of input folder where coded sheets are stored
#(With reference to parent program folder)
input_directory = "Coded"

###############################################

mstr_wb = openpyxl.Workbook()
phrase_ws = mstr_wb.active

#initialize sheets
phrase_ws.title = "PHRASES"
phrase_ws['D1'].value = "PHRASES: TOTAL WEIGHT"
phrase_ws['G1'].value = "PHRASES: RELATIVE WEIGHT"

word_ws = mstr_wb.create_sheet(title="WORDS",index=1)
word_ws['D1'].value = "WORDS: TOTAL WEIGHT"
word_ws['G1'].value = "WORDS: RELATIVE WEIGHT"

all_ws = mstr_wb.create_sheet(title="WORDS&PHRASES",index=2)
all_ws['D1'].value = "WORDS&PHRASES: TOTAL WEIGHT"
all_ws['G1'].value = "WORDS&PHRASES: RELATIVE WEIGHT"

#initialize master row
mstr_row = 3

def construct_format(mstr_ws):
    #label categories for table
    mstr_ws['C1'].fill = openpyxl.styles.fills.PatternFill(start_color="BEBEBE", end_color="BEBEBE", fill_type="solid")
    mstr_ws['D1'].fill = openpyxl.styles.fills.PatternFill(start_color="BEBEBE", end_color="BEBEBE", fill_type="solid")
    mstr_ws['E1'].fill = openpyxl.styles.fills.PatternFill(start_color="BEBEBE", end_color="BEBEBE", fill_type="solid")
    mstr_ws['F1'].fill = openpyxl.styles.fills.PatternFill(start_color="686868", end_color="686868", fill_type="solid")
    mstr_ws['G1'].fill = openpyxl.styles.fills.PatternFill(start_color="686868", end_color="686868", fill_type="solid")
    mstr_ws['H1'].fill = openpyxl.styles.fills.PatternFill(start_color="686868", end_color="686868", fill_type="solid")
    
    #label table
    mstr_ws['A2'] = "Country"
    mstr_ws['B2'] = "Strategy"
    mstr_ws['C2'] = "Deter"
    mstr_ws['D2'] = "Norms"
    mstr_ws['E2'] = "PE"
    mstr_ws['F2'] = "Deter"
    mstr_ws['G2'] = "Norms"
    mstr_ws['H2'] = "PE"

    #highlight table
    #green (deterrence)
    mstr_ws['C2'].fill = openpyxl.styles.fills.PatternFill(start_color="E6FFCC", end_color="E6FFCC", fill_type="solid")
    mstr_ws['F2'].fill = openpyxl.styles.fills.PatternFill(start_color="E6FFCC", end_color="E6FFCC", fill_type="solid")
    #blue (norms)
    mstr_ws['D2'].fill = openpyxl.styles.fills.PatternFill(start_color="99CCFF", end_color="99CCFF", fill_type="solid")
    mstr_ws['G2'].fill = openpyxl.styles.fills.PatternFill(start_color="99CCFF", end_color="99CCFF", fill_type="solid")
    #purple (PE)
    mstr_ws['E2'].fill = openpyxl.styles.fills.PatternFill(start_color="B3B3FF", end_color="B3B3FF", fill_type="solid")
    mstr_ws['H2'].fill = openpyxl.styles.fills.PatternFill(start_color="B3B3FF", end_color="B3B3FF", fill_type="solid")

def get_country(file_name):
    country = ""
    if file_name.startswith("new_"):
        for i in range(4, len(file_name)):
            if file_name[i].isalpha():
                country += file_name[i]
            else:
                return country
    else:
        print("file naming error")

def get_year(file_name):
    year = ""
    if file_name.startswith("new_"):
        for i in range(4, len(file_name)):
            if file_name[i].isdigit():
                year += file_name[i]
    else:
        print("file naming error")
    
    return year



def collect_data(file_name):
    country = get_country(file_name)
    year = get_year(file_name)
    
    tracking_cols = ['G','H','P','Q','Y','Z', 'G', 'G']
    i = 0

    for sheet in mstr_wb:
        #get country and year
        sheet.cell(row=mstr_row, column=1).value = country
        sheet.cell(row=mstr_row, column=2).value = year

        #total values
        sheet.cell(row=mstr_row, column=3).value = ws[tracking_cols[i] + '2'].value
        sheet.cell(row=mstr_row, column=4).value = ws[tracking_cols[i] + '3'].value
        sheet.cell(row=mstr_row, column=5).value = ws[tracking_cols[i] + '4'].value

        i += 1

        #relative values
        sheet.cell(row=mstr_row, column=6).value = ws[tracking_cols[i] + '2'].value
        sheet.cell(row=mstr_row, column=7).value = ws[tracking_cols[i] + '3'].value
        sheet.cell(row=mstr_row, column=8).value = ws[tracking_cols[i] + '4'].value

        i += 1
    


for sheet in mstr_wb:
    construct_format(sheet)

for entry in os.scandir(input_directory):
    if entry.name.endswith(".xlsx"):
        wb = xw.Book(entry)
        ws = wb.sheets[0]
    
    collect_data(entry.name)
    mstr_row += 1

    wb.close()

mstr_wb.save(filename="Mastersheet.xlsx")